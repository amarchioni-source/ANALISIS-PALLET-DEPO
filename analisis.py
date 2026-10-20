"""Lógica de cálculo: stock por pallet (mono / multiproducto)."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

COLS = ["PALLET", "CODIGO", "CAJAS", "kg", "camara"]
RANGOS_CAJAS = [("Mayor a 40 cajas", lambda s: s > 40),
                ("Entre 30 y 40 cajas", lambda s: (s >= 30) & (s <= 40)),
                ("Menor a 30 cajas", lambda s: s < 30)]
RANGOS_KG = [("Mayor a 800 kg", lambda s: s > 800),
             ("Entre 600 - 800 kg", lambda s: (s >= 600) & (s <= 800)),
             ("Menor a 600 kg", lambda s: s < 600)]


def _leer(fileobj):
    try:
        engine = "calamine"
        import python_calamine  # noqa: F401
    except ImportError:
        engine = "openpyxl"
    raw = pd.read_excel(fileobj, header=None, nrows=15, engine=engine)
    fila = next((i for i in range(len(raw)) if "PALLET" in raw.iloc[i].astype(str).values), None)
    if fila is None:
        raise ValueError("No se encontró la columna PALLET en las primeras filas del archivo.")
    fileobj.seek(0)
    df = pd.read_excel(fileobj, header=fila, engine=engine, usecols=lambda c: str(c).strip() in COLS)
    df.columns = [str(c).strip() for c in df.columns]
    faltan = [c for c in COLS if c not in df.columns]
    if faltan:
        raise ValueError(f"Faltan columnas en el archivo: {', '.join(faltan)}")
    return df


def calcular(fileobj):
    df = _leer(fileobj)
    total_cajas_archivo = len(df)
    p = df[df["PALLET"].notna()].copy()
    p["CAJAS"] = pd.to_numeric(p["CAJAS"], errors="coerce").fillna(0)
    p["kg"] = pd.to_numeric(p["kg"], errors="coerce").fillna(0)
    g = (p.groupby("PALLET")
           .agg(camara=("camara", "first"), productos=("CODIGO", "nunique"),
                cajas=("CAJAS", "sum"), kg=("kg", "sum"))
           .reset_index())
    g["tipo"] = g["productos"].map(lambda x: "Mono" if x == 1 else "Multi")
    total = len(g)

    def bloque(rangos, col):
        out = {"total": total}
        for tipo in ("Mono", "Multi"):
            sub = g[g["tipo"] == tipo]
            n = len(sub)
            filas = []
            for nombre, f in rangos:
                c = int(f(sub[col]).sum())
                filas.append({"label": nombre, "n": c, "pct": c / n if n else 0})
            out[tipo] = {"n": n, "pct": n / total if total else 0, "filas": filas}
        return out

    return {
        "cajas": bloque(RANGOS_CAJAS, "cajas"),
        "kg": bloque(RANGOS_KG, "kg"),
        "sin_pallet": total_cajas_archivo - len(p),
        "detalle": g,
    }


def generar_excel(res, ruta):
    g = res["detalle"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    d = wb.create_sheet("Detalle_Pallets")
    F = "Arial"
    for i, h in enumerate(["Pallet", "Cámara", "Productos distintos", "Tipo", "Cajas", "Kg"], 1):
        c = d.cell(1, i, h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="000000")
    for r, row in enumerate(g.itertuples(index=False), 2):
        d.cell(r, 1, row.PALLET); d.cell(r, 2, row.camara); d.cell(r, 3, int(row.productos))
        d.cell(r, 4, f'=IF(C{r}=1,"Mono","Multi")')
        d.cell(r, 5, int(row.cajas)); d.cell(r, 6, round(float(row.kg), 2))
        for i in range(1, 7):
            d.cell(r, i).font = Font(name=F)
    n = len(g) + 1
    for col, w in zip("ABCDEF", [20, 24, 20, 10, 10, 12]):
        d.column_dimensions[col].width = w
    d.freeze_panes = "A2"
    d.auto_filter.ref = f"A1:F{n}"
    D = lambda c: f"Detalle_Pallets!${c}$2:${c}${n}"

    def block(r0, labels, conds):
        ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=6)
        c = ws.cell(r0, 1, f'="Pallets en stock   "&COUNTA({D("A")})')
        c.font = Font(name=F, bold=True, size=14, color="FFFFFF")
        c.alignment = Alignment(horizontal="center")
        for i in range(1, 7):
            ws.cell(r0, i).fill = PatternFill("solid", fgColor="000000")
        tot = f'COUNTA({D("A")})'
        for side, (name, tipo, fill) in enumerate([("Mono Producto", "Mono", "C6EFCE"), ("Multiproducto", "Multi", "F2CEEF")]):
            c0 = 1 + side * 3
            L = chr(64 + c0 + 1)
            ws.cell(r0 + 1, c0, name)
            ws.cell(r0 + 1, c0 + 1, f'=COUNTIFS({D("D")},"{tipo}")')
            ws.cell(r0 + 1, c0 + 2, f"=IFERROR({L}{r0+1}/{tot},0)")
            for k in range(3):
                x = ws.cell(r0 + 1, c0 + k)
                x.fill = PatternFill("solid", fgColor=fill)
                x.font = Font(name=F, bold=True, size=14)
                x.alignment = Alignment(horizontal="left" if k == 0 else "center")
            ws.cell(r0 + 1, c0 + 2).number_format = "0%"
            for j, (lab, crit, color) in enumerate(zip(labels, conds, ["2E75B6", "548235", "FF0000"])):
                r = r0 + 2 + j
                ws.cell(r, c0, lab)
                ws.cell(r, c0 + 1, f'=COUNTIFS({D("D")},"{tipo}",{crit})')
                ws.cell(r, c0 + 2, f"=IFERROR({L}{r}/{L}{r0+1},0)")
                ws.cell(r, c0 + 2).number_format = "0%"
                for k in range(3):
                    x = ws.cell(r, c0 + k)
                    x.fill = PatternFill("solid", fgColor="D9D9D9")
                    x.font = Font(name=F, bold=True, size=14, color=color)
                    x.alignment = Alignment(horizontal="left" if k == 0 else "center")

    cajas = [f'{D("E")},">40"', f'{D("E")},">=30",{D("E")},"<=40"', f'{D("E")},"<30"']
    kgs = [f'{D("F")},">800"', f'{D("F")},">=600",{D("F")},"<=800"', f'{D("F")},"<600"']
    block(1, [r[0] for r in RANGOS_CAJAS], cajas)
    block(7, [r[0] for r in RANGOS_KG], kgs)
    for col, w in zip("ABCDEF", [26, 10, 8, 26, 10, 8]):
        ws.column_dimensions[col].width = w
    ws.cell(13, 1, 'Pallet = valor único de la columna PALLET. Mono = 1 código de producto; Multi = 2 o más. '
                   'Se excluyen cajas sin pallet. Los límites (30/40 cajas, 600/800 kg) están incluidos en "Entre".'
            ).font = Font(name=F, italic=True, size=9)
    wb.save(ruta)
